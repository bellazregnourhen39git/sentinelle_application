import docx
import json

doc = docx.Document('MedSPAD_Questionnaire Français 2026 (1).docx')
paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

intro_paras = []
for i, p in enumerate(paras):
    if any(k in p.lower() for k in ['premières questions', 'questions suivantes concernent', 'questions suivantes portent', 'questions portent sur', 'concernent vos parents', 'questions finales', 'questions suivantes ont trait', 'questions suivantes portent sur']):
        intro_paras.append((i, p))
    elif p.startswith('Les questions') or p.startswith('Les premières questions'):
        intro_paras.append((i, p))

print(f"Found {len(intro_paras)} intro paragraphs:")
for idx, text in intro_paras:
    print(f"[{idx}] {text}\n")

# Also check tables for intro text
for t_idx, table in enumerate(doc.tables):
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            t = cell.text.strip()
            if any(k in t.lower() for k in ['premières questions', 'questions suivantes concernent', 'concernent vos parents', 'questions finales', 'sans compter les cigarettes']):
                if len(t) < 500:
                    print(f"Table {t_idx} [r{r_idx},c{c_idx}]: {t}\n")
